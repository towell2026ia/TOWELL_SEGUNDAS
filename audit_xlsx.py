import openpyxl,json,collections,datetime,pathlib,hashlib
source=pathlib.Path(r'C:\Users\gcalidad\OneDrive - Towell\Documentos\AX\Segundas 2026.xlsx')
w=openpyxl.load_workbook(source,read_only=True,data_only=True)
out={}
for name in ['PD','%Segundas','SEGUNDAS','COSTURA','MAQUILA','Segundas Crudo']:
 s=w[name]; rows=[]; totals=collections.defaultdict(float); dates=[]; counts=collections.Counter()
 for i,r in enumerate(s.iter_rows(values_only=True),1):
  if name=='PD' or (name=='%Segundas' and i<=180):
   rows.append([i,[(openpyxl.utils.get_column_letter(j+1),str(x)) for j,x in enumerate(r) if x is not None]])
  elif i<=5: rows.append([i,[str(x) if x is not None else None for x in r]])
  if name=='%Segundas' and isinstance(r[0],datetime.datetime):
   totals['production_table']+=r[6] or 0; totals[str(r[1])]+=r[6] or 0;dates.append(r[0]);counts['production_rows']+=1
  if name=='SEGUNDAS' and i>1 and isinstance(r[5],datetime.datetime):
   dates.append(r[5]);totals['segundas_costura']+=r[15] or 0;totals['total_segundas']+=r[19] or 0;counts[str(r[0])]+=1
  if name=='COSTURA' and i>1 and isinstance(r[0],datetime.datetime):
   dates.append(r[0]);totals['quantity']+=r[11] or 0;totals[str(r[9])]+=r[11] or 0
  if name=='MAQUILA' and i>1 and isinstance(r[1],datetime.datetime):
   dates.append(r[1]);totals['InventQty']+=r[12] or 0;totals['Consumo']+=r[26] or 0;totals['AE']+=r[30] or 0
  if name=='Segundas Crudo' and i>1 and isinstance(r[1],datetime.datetime):
   dates.append(r[1]);totals['quantity']+=r[29] or 0
 out[name]={'rows':rows,'totals':dict(totals),'counts':dict(counts),'date_range':[str(min(dates)),str(max(dates))] if dates else []}
pathlib.Path('audit-xlsx.json').write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:{a:b for a,b in v.items() if a!='rows'} for k,v in out.items()},ensure_ascii=True))
w.close()
wf=openpyxl.load_workbook(source,read_only=True,data_only=False)
for name,area in [('MAQUILA',(1,8,29,37)),('%Segundas',(8,25,13,16)),('PD',(1,203,1,17))]:
 s=wf[name];print(name,'FORMULAS',[(c.coordinate,c.value) for row in s.iter_rows(min_row=area[0],max_row=area[1],min_col=area[2],max_col=area[3]) for c in row if c.data_type=='f'][:80])
